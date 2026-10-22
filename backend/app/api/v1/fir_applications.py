"""
FIR VisionAI — FIR Application Management API
Handles multilingual form submission, date-wise search, and export.
All endpoints enforce jurisdiction-based access control.
"""
from __future__ import annotations
import io
import csv
import json
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from datetime import datetime
from typing import Optional
from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import Response
from pydantic import BaseModel
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy import select, func, or_
from ...database import get_db
from ...models.case import Case
from ...models.user import User
from ...services.auth import get_current_user, apply_case_scope_filter, enforce_data_scope
from ...services.audit import log_audit_event

router = APIRouter(prefix="/fir-applications", tags=["FIR Applications"])


# ─── Pydantic schemas ────────────────────────────────────────────────────────

class FIRApplicationCreate(BaseModel):
    fir_number: Optional[str] = None
    district: str
    station: str
    mandal: Optional[str] = None
    village: Optional[str] = None
    location_description: Optional[str] = None
    offence_nature: str = "FIR Application"
    offence_category: Optional[str] = None
    contraband_type: Optional[str] = None
    seized_quantity: Optional[float] = None
    seized_unit: Optional[str] = None
    vehicle_details: Optional[str] = None
    seizing_officer_name: Optional[str] = None
    seizing_officer_rank: Optional[str] = None
    witnesses_summary: Optional[str] = None
    narrative: Optional[str] = None
    application_date: Optional[str] = None
    incident_date: Optional[str] = None
    form_template: Optional[str] = "FIR_FORM_AP_EXCISE_v1"
    form_language: Optional[str] = "en"
    form_values_json: Optional[str] = None
    application_status: Optional[str] = "DRAFT"


class FIRApplicationUpdate(BaseModel):
    mandal: Optional[str] = None
    village: Optional[str] = None
    location_description: Optional[str] = None
    offence_nature: Optional[str] = None
    offence_category: Optional[str] = None
    contraband_type: Optional[str] = None
    seized_quantity: Optional[float] = None
    seized_unit: Optional[str] = None
    vehicle_details: Optional[str] = None
    seizing_officer_name: Optional[str] = None
    seizing_officer_rank: Optional[str] = None
    witnesses_summary: Optional[str] = None
    narrative: Optional[str] = None
    incident_date: Optional[str] = None
    form_language: Optional[str] = None
    form_values_json: Optional[str] = None
    application_status: Optional[str] = None


# ─── Helpers ─────────────────────────────────────────────────────────────────

def _parse_date(val: Optional[str]) -> Optional[datetime]:
    if not val:
        return None
    for fmt in ("%Y-%m-%d", "%Y-%m-%dT%H:%M:%S", "%Y-%m-%dT%H:%M:%SZ", "%d/%m/%Y"):
        try:
            return datetime.strptime(val.strip(), fmt)
        except ValueError:
            continue
    return None


def _case_dict(c: Case) -> dict:
    return {
        "id": c.id,
        "fir_number": c.fir_number,
        "year": c.year,
        "district": c.district,
        "station": c.station,
        "mandal": c.mandal,
        "village": c.village,
        "location_description": c.location_description,
        "offence_nature": c.offence_nature,
        "offence_category": c.offence_category,
        "contraband_type": c.contraband_type,
        "seized_quantity": c.seized_quantity,
        "seized_unit": c.seized_unit,
        "vehicle_details": c.vehicle_details,
        "seizing_officer_name": c.seizing_officer_name,
        "seizing_officer_rank": c.seizing_officer_rank,
        "witnesses_summary": c.witnesses_summary,
        "narrative": c.narrative,
        "status": c.status,
        "application_status": getattr(c, "application_status", None),
        "form_template": getattr(c, "form_template", None),
        "form_language": getattr(c, "form_language", None),
        "form_values_json": getattr(c, "form_values_json", None),
        "application_date": c.application_date.isoformat() if getattr(c, "application_date", None) else None,
        "incident_date": c.incident_date.isoformat() if getattr(c, "incident_date", None) else None,
        "registration_date": c.registration_date.isoformat() if c.registration_date else None,
        "offence_date": c.offence_date.isoformat() if c.offence_date else None,
        "created_at": c.created_at.isoformat() if c.created_at else None,
        "updated_at": c.updated_at.isoformat() if c.updated_at else None,
    }


def _apply_date_filter(query, date_type: str, start: Optional[datetime], end: Optional[datetime]):
    if not start and not end:
        return query
    if date_type == "registration":
        col = Case.registration_date
    elif date_type == "incident":
        col = Case.incident_date
    else:
        col = Case.application_date
    if start:
        query = query.where(col >= start.replace(hour=0, minute=0, second=0))
    if end:
        query = query.where(col <= end.replace(hour=23, minute=59, second=59))
    return query


# ─── Endpoints ───────────────────────────────────────────────────────────────

@router.get("/summary")
async def get_daily_summary(
    date: Optional[str] = Query(None),
    date_type: str = Query("application"),
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    """Daily case counts for a specific date."""
    target = _parse_date(date) if date else datetime.utcnow()
    if not target:
        target = datetime.utcnow()
    start = target.replace(hour=0, minute=0, second=0, microsecond=0)
    end = target.replace(hour=23, minute=59, second=59, microsecond=999999)

    base = select(Case).where(Case.is_archived == False)
    base = apply_case_scope_filter(base, current_user, Case)
    base = _apply_date_filter(base, date_type, start, end)

    total = (await db.execute(select(func.count()).select_from(base.subquery()))).scalar() or 0
    registered = (await db.execute(
        select(func.count()).select_from(
            base.where(or_(Case.status == "APPROVED", Case.application_status == "REGISTERED")).subquery()
        )
    )).scalar() or 0
    drafts = (await db.execute(
        select(func.count()).select_from(
            base.where(or_(Case.application_status == "DRAFT", Case.status == "DRAFT")).subquery()
        )
    )).scalar() or 0
    under_inv = (await db.execute(
        select(func.count()).select_from(
            base.where(or_(Case.application_status == "UNDER_INVESTIGATION", Case.status == "IN_REVIEW")).subquery()
        )
    )).scalar() or 0

    return {
        "date": target.strftime("%Y-%m-%d"),
        "date_type": date_type,
        "total_applications": total,
        "registered": registered,
        "drafts": drafts,
        "under_investigation": under_inv,
    }


@router.get("/export/results")
async def export_search_results(
    fmt: str = Query("xlsx"),
    date: Optional[str] = None,
    from_date: Optional[str] = None,
    to_date: Optional[str] = None,
    date_type: str = Query("application"),
    fir_number: Optional[str] = None,
    status: Optional[str] = None,
    application_status: Optional[str] = None,
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    """Export date-filtered results as xlsx, csv, or json."""
    query = select(Case).where(Case.is_archived == False)
    query = apply_case_scope_filter(query, current_user, Case)

    if date:
        dt = _parse_date(date)
        query = _apply_date_filter(query, date_type, dt, dt)
    else:
        query = _apply_date_filter(query, date_type, _parse_date(from_date), _parse_date(to_date))

    if fir_number:
        query = query.where(Case.fir_number.ilike(f"%{fir_number}%"))
    if status:
        query = query.where(Case.status == status.upper())
    if application_status:
        query = query.where(Case.application_status == application_status.upper())

    query = query.order_by(Case.application_date.desc().nullslast(), Case.created_at.desc()).limit(1000)
    res = await db.execute(query)
    cases = res.scalars().all()
    rows = [_case_dict(c) for c in cases]
    ts = datetime.utcnow().strftime("%Y%m%d_%H%M%S")

    if fmt == "json":
        content = json.dumps(rows, ensure_ascii=False, indent=2).encode("utf-8")
        media_type = "application/json; charset=utf-8"
        filename = f"FIR_Applications_{ts}.json"
    elif fmt == "csv":
        buf = io.StringIO()
        if rows:
            writer = csv.DictWriter(buf, fieldnames=list(rows[0].keys()))
            writer.writeheader()
            writer.writerows(rows)
        content = buf.getvalue().encode("utf-8-sig")
        media_type = "text/csv; charset=utf-8"
        filename = f"FIR_Applications_{ts}.csv"
    elif fmt == "xlsx":
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "FIR Applications"
        if rows:
            headers = list(rows[0].keys())
            hf = Font(bold=True, color="FFFFFF")
            hfill = PatternFill(fill_type="solid", fgColor="0F5132")
            for ci, h in enumerate(headers, 1):
                cell = ws.cell(row=1, column=ci, value=h.replace("_", " ").title())
                cell.font = hf
                cell.fill = hfill
                cell.alignment = Alignment(horizontal="center")
            for ri, row in enumerate(rows, 2):
                for ci, key in enumerate(headers, 1):
                    val = row.get(key)
                    ws.cell(row=ri, column=ci, value=str(val) if val is not None else "")
            for col in ws.columns:
                max_l = max((len(str(c.value or "")) for c in col), default=10)
                ws.column_dimensions[col[0].column_letter].width = min(max_l + 2, 40)
        buf = io.BytesIO()
        wb.save(buf)
        content = buf.getvalue()
        media_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        filename = f"FIR_Applications_{ts}.xlsx"
    else:
        raise HTTPException(status_code=400, detail="Unsupported format. Use xlsx, csv, or json.")

    await log_audit_event(
        db=db, actor_id=current_user.id, actor_username=current_user.username,
        action_type="FIR_APPLICATIONS_EXPORTED", object_type="CASE_LIST", object_id="bulk",
        after_state={"format": fmt, "count": len(rows)},
    )
    return Response(content=content, media_type=media_type,
                    headers={"Content-Disposition": f'attachment; filename="{filename}"',
                             "Access-Control-Expose-Headers": "Content-Disposition"})


@router.get("/{case_id}/download")
async def download_fir_form(
    case_id: str,
    fmt: str = Query("pdf"),
    lang: str = Query("en"),
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    """Download an individual FIR application as pdf, docx, or json."""
    res = await db.execute(select(Case).where(Case.id == case_id))
    case = res.scalar_one_or_none()
    if not case:
        raise HTTPException(status_code=404, detail="FIR application not found")
    if not enforce_data_scope(current_user, case.district, case.station):
        raise HTTPException(status_code=403, detail="Access denied: outside your jurisdiction")

    from ...models.document import Document
    from ...models.case import CaseField
    from ...models.folder import WorkflowExecution

    f_res = await db.execute(select(CaseField).where(CaseField.case_id == case_id))
    fields = [{"field_name": f.field_name, "normalized_value": f.normalized_value,
                "confidence_score": f.confidence_score, "verification_status": f.verification_status}
               for f in f_res.scalars().all()]

    w_res = await db.execute(select(WorkflowExecution).where(WorkflowExecution.case_id == case_id))
    workflow = [{"stage": w.stage, "status": w.status,
                 "started_at": w.started_at.isoformat() if w.started_at else None,
                 "completed_at": w.completed_at.isoformat() if w.completed_at else None}
                for w in w_res.scalars().all()]

    export_data = {
        "case": _case_dict(case),
        "fields": fields,
        "entities": [],
        "documents": [],
        "workflow": workflow,
        "language": lang,
    }

    from ...services.case_export import generate_case_pdf, generate_case_docx, generate_case_json
    clean_fir = (case.fir_number or "FIR_Application").replace("/", "_").replace("\\", "_")
    ts = datetime.utcnow().strftime("%Y%m%d_%H%M%S")
    filename = f"{clean_fir}_{ts}.{fmt}"

    try:
        if fmt == "pdf":
            content = generate_case_pdf(export_data)
            media_type = "application/pdf"
        elif fmt == "docx":
            content = generate_case_docx(export_data)
            media_type = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        elif fmt == "json":
            content = generate_case_json(export_data)
            media_type = "application/json; charset=utf-8"
        else:
            raise HTTPException(status_code=400, detail="Use pdf, docx, or json.")
    except Exception as exc:
        raise HTTPException(status_code=500, detail=f"Export failed: {str(exc)}")

    await log_audit_event(
        db=db, actor_id=current_user.id, actor_username=current_user.username,
        action_type="FIR_FORM_DOWNLOADED", object_type="CASE", object_id=case_id,
        after_state={"format": fmt, "language": lang},
    )
    return Response(content=content, media_type=media_type,
                    headers={"Content-Disposition": f'attachment; filename="{filename}"',
                             "Access-Control-Expose-Headers": "Content-Disposition"})


@router.get("/")
async def search_fir_applications(
    date: Optional[str] = Query(None),
    from_date: Optional[str] = Query(None),
    to_date: Optional[str] = Query(None),
    date_type: str = Query("application"),
    fir_number: Optional[str] = None,
    case_id: Optional[str] = None,
    person: Optional[str] = None,
    location: Optional[str] = None,
    status: Optional[str] = None,
    application_status: Optional[str] = None,
    form_language: Optional[str] = None,
    form_template: Optional[str] = None,
    page: int = 1,
    limit: int = 30,
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    """Date-wise search of FIR applications."""
    query = select(Case).where(Case.is_archived == False)
    query = apply_case_scope_filter(query, current_user, Case)

    if date:
        dt = _parse_date(date)
        query = _apply_date_filter(query, date_type, dt, dt)
    else:
        query = _apply_date_filter(query, date_type, _parse_date(from_date), _parse_date(to_date))

    if fir_number:
        query = query.where(Case.fir_number.ilike(f"%{fir_number}%"))
    if case_id:
        query = query.where(Case.id == case_id)
    if location:
        query = query.where(or_(
            Case.mandal.ilike(f"%{location}%"),
            Case.village.ilike(f"%{location}%"),
            Case.location_description.ilike(f"%{location}%"),
        ))
    if person:
        query = query.where(or_(
            Case.seizing_officer_name.ilike(f"%{person}%"),
            Case.narrative.ilike(f"%{person}%"),
        ))
    if status:
        query = query.where(Case.status == status.upper())
    if application_status:
        query = query.where(Case.application_status == application_status.upper())
    if form_language:
        query = query.where(Case.form_language == form_language)
    if form_template:
        query = query.where(Case.form_template == form_template)

    total = (await db.execute(select(func.count()).select_from(query.subquery()))).scalar() or 0
    query = query.order_by(Case.application_date.desc().nullslast(), Case.created_at.desc())
    query = query.offset((page - 1) * limit).limit(limit)
    res = await db.execute(query)
    cases = res.scalars().all()

    return {"total": total, "page": page, "limit": limit, "cases": [_case_dict(c) for c in cases]}


@router.get("/{case_id}")
async def get_fir_application(
    case_id: str,
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    res = await db.execute(select(Case).where(Case.id == case_id))
    case = res.scalar_one_or_none()
    if not case:
        raise HTTPException(status_code=404, detail="FIR application not found")
    if not enforce_data_scope(current_user, case.district, case.station):
        raise HTTPException(status_code=403, detail="Access denied")
    return _case_dict(case)


@router.post("/")
async def create_fir_application(
    body: FIRApplicationCreate,
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    """Create a new FIR application draft or submission."""
    from ...models.base import generate_uuid

    if not enforce_data_scope(current_user, body.district, body.station):
        raise HTTPException(status_code=403, detail="Cannot create application outside your jurisdiction")

    if body.fir_number:
        existing = (await db.execute(select(Case).where(Case.fir_number == body.fir_number))).scalar_one_or_none()
        if existing:
            raise HTTPException(status_code=409, detail=f"Case with FIR number '{body.fir_number}' already exists")

    now = datetime.utcnow()
    fir_number = body.fir_number or f"DRAFT-{now.strftime('%Y%m%d%H%M%S')}-{generate_uuid()[:8].upper()}"

    case = Case(
        id=generate_uuid(),
        fir_number=fir_number,
        year=now.year,
        district=body.district,
        station=body.station,
        mandal=body.mandal,
        village=body.village,
        location_description=body.location_description,
        offence_nature=body.offence_nature or "FIR Application",
        offence_category=body.offence_category,
        contraband_type=body.contraband_type,
        seized_quantity=body.seized_quantity,
        seized_unit=body.seized_unit,
        vehicle_details=body.vehicle_details,
        seizing_officer_name=body.seizing_officer_name,
        seizing_officer_rank=body.seizing_officer_rank,
        witnesses_summary=body.witnesses_summary,
        narrative=body.narrative,
        status="DRAFT",
        is_archived=False,
        created_by=current_user.id,
        registration_date=now,
        application_date=_parse_date(body.application_date) or now,
        incident_date=_parse_date(body.incident_date),
        form_template=body.form_template or "FIR_FORM_AP_EXCISE_v1",
        form_language=body.form_language or "en",
        form_values_json=body.form_values_json,
        application_status=body.application_status or "DRAFT",
        classification_status="PENDING",
    )
    db.add(case)
    await db.commit()
    await db.refresh(case)

    await log_audit_event(
        db=db, actor_id=current_user.id, actor_username=current_user.username,
        action_type="FIR_APPLICATION_CREATED", object_type="CASE", object_id=case.id,
        after_state={"fir_number": fir_number, "application_status": case.application_status},
    )
    return _case_dict(case)


@router.put("/{case_id}")
async def update_fir_application(
    case_id: str,
    body: FIRApplicationUpdate,
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    res = await db.execute(select(Case).where(Case.id == case_id))
    case = res.scalar_one_or_none()
    if not case:
        raise HTTPException(status_code=404, detail="FIR application not found")
    if not enforce_data_scope(current_user, case.district, case.station):
        raise HTTPException(status_code=403, detail="Access denied")

    updates = body.dict(exclude_none=True)
    for field, value in updates.items():
        if field == "incident_date":
            setattr(case, field, _parse_date(value))
        elif hasattr(case, field):
            setattr(case, field, value)

    await db.commit()
    await db.refresh(case)

    await log_audit_event(
        db=db, actor_id=current_user.id, actor_username=current_user.username,
        action_type="FIR_APPLICATION_UPDATED", object_type="CASE", object_id=case_id,
        after_state={"application_status": case.application_status},
    )
    return _case_dict(case)
