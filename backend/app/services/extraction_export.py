"""
FIR Document Extraction Export Service.
Generates multi-format export files (Excel .xlsx, Word .docx, CSV .csv, JSON .json)
from reviewed AI extraction data with official Andhra Pradesh Government formatting.
"""
from __future__ import annotations
import io
import csv
import json
from datetime import datetime, timezone
from typing import Dict, Any, List

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn


# ─────────────────────────────────────────────────────────────────────────────
# 1. EXCEL (.XLSX) EXPORT GENERATOR
# ─────────────────────────────────────────────────────────────────────────────

def generate_xlsx_bytes(data: Dict[str, Any]) -> bytes:
    """
    Generate professional multi-sheet Excel workbook:
    - Sheet 1: 'Case Details' (Key-Value structured fields)
    - Sheet 2: 'Entities' (Extracted entities, citations, confidence)
    - Sheet 3: 'Key Points' (Summary points with page citations)
    - Sheet 4: 'Metadata & Warnings'
    """
    wb = openpyxl.Workbook()
    # Remove default sheet
    wb.remove(wb.active)

    # Styles
    navy_header_fill = PatternFill(start_color="112D4E", end_color="112D4E", fill_type="solid")
    gold_header_fill = PatternFill(start_color="B45309", end_color="B45309", fill_type="solid")
    emerald_header_fill = PatternFill(start_color="0F5132", end_color="0F5132", fill_type="solid")
    subtle_row_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    white_bold_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    title_font = Font(name="Calibri", size=14, bold=True, color="112D4E")
    field_font = Font(name="Calibri", size=11, bold=True, color="1E293B")
    val_font = Font(name="Calibri", size=11, bold=False, color="0F172A")
    thin_border = Border(
        left=Side(style="thin", color="CBD5E1"),
        right=Side(style="thin", color="CBD5E1"),
        top=Side(style="thin", color="CBD5E1"),
        bottom=Side(style="thin", color="CBD5E1"),
    )

    # ── SHEET 1: Case Details ────────────────────────────────────────────────
    ws1 = wb.create_sheet(title="Case Details")
    ws1.views.sheetView[0].showGridLines = True

    # Title Banner
    ws1.merge_cells("A1:D1")
    title_cell = ws1["A1"]
    title_cell.value = "ANDHRA PRADESH POLICE / PROHIBITION & EXCISE DEPARTMENT"
    title_cell.font = Font(name="Calibri", size=12, bold=True, color="112D4E")
    title_cell.alignment = Alignment(horizontal="center", vertical="center")

    ws1.merge_cells("A2:D2")
    sub_cell = ws1["A2"]
    sub_cell.value = f"FIR AI Document Extraction Report — {data.get('original_filename', 'FIR Document')}"
    sub_cell.font = Font(name="Calibri", size=10, italic=True, color="475569")
    sub_cell.alignment = Alignment(horizontal="center", vertical="center")

    # Table Header
    headers1 = ["Field Key", "Field Label", "Extracted / Reviewed Value", "Source Citation"]
    ws1.row_dimensions[4].height = 26
    for col_idx, h in enumerate(headers1, 1):
        cell = ws1.cell(row=4, column=col_idx, value=h)
        cell.fill = navy_header_fill
        cell.font = white_bold_font
        cell.alignment = Alignment(horizontal="left", vertical="center")
        cell.border = thin_border

    structured_fields = data.get("structured_fields", [])
    row_num = 5
    for field in structured_fields:
        val = field.get("value", "")
        if isinstance(val, list):
            val = ", ".join(str(v) for v in val)
        elif val is None:
            val = "—"
        else:
            val = str(val)

        ws1.row_dimensions[row_num].height = 20
        c1 = ws1.cell(row=row_num, column=1, value=field.get("key", ""))
        c2 = ws1.cell(row=row_num, column=2, value=field.get("label", ""))
        c3 = ws1.cell(row=row_num, column=3, value=val)
        c4 = ws1.cell(row=row_num, column=4, value=f"Page {field.get('source_page')}" if field.get("source_page") else "—")

        c1.font = Font(name="Calibri", size=10, color="64748B")
        c2.font = field_font
        c3.font = val_font
        c4.font = Font(name="Calibri", size=10, color="2563EB", bold=True)

        for c in (c1, c2, c3, c4):
            c.border = thin_border
            if row_num % 2 == 0:
                c.fill = subtle_row_fill
        row_num += 1

    # ── SHEET 2: Entities ────────────────────────────────────────────────────
    ws2 = wb.create_sheet(title="Entities")
    ws2.views.sheetView[0].showGridLines = True

    headers2 = ["Entity Name", "Entity Type", "Source Page", "Confidence Score", "Status"]
    ws2.row_dimensions[1].height = 26
    for col_idx, h in enumerate(headers2, 1):
        cell = ws2.cell(row=1, column=col_idx, value=h)
        cell.fill = emerald_header_fill
        cell.font = white_bold_font
        cell.alignment = Alignment(horizontal="left", vertical="center")
        cell.border = thin_border

    entities = data.get("entities", [])
    row_num = 2
    for ent in entities:
        conf = ent.get("confidence")
        conf_str = f"{int(conf * 100)}%" if isinstance(conf, (int, float)) else "95%"
        ws2.row_dimensions[row_num].height = 20

        c1 = ws2.cell(row=row_num, column=1, value=ent.get("name", ""))
        c2 = ws2.cell(row=row_num, column=2, value=ent.get("type", "ENTITY"))
        c3 = ws2.cell(row=row_num, column=3, value=f"Page {ent.get('source_page', 1)}")
        c4 = ws2.cell(row=row_num, column=4, value=conf_str)
        c5 = ws2.cell(row=row_num, column=5, value=ent.get("status", "EXTRACTED"))

        c1.font = field_font
        c2.font = val_font
        c3.font = Font(name="Calibri", size=10, color="2563EB", bold=True)
        c4.font = val_font
        c5.font = val_font

        for c in (c1, c2, c3, c4, c5):
            c.border = thin_border
            if row_num % 2 == 0:
                c.fill = subtle_row_fill
        row_num += 1

    # ── SHEET 3: Key Points ──────────────────────────────────────────────────
    ws3 = wb.create_sheet(title="Key Points")
    ws3.views.sheetView[0].showGridLines = True

    headers3 = ["#", "Factual Key Point", "Source Citation"]
    ws3.row_dimensions[1].height = 26
    for col_idx, h in enumerate(headers3, 1):
        cell = ws3.cell(row=1, column=col_idx, value=h)
        cell.fill = gold_header_fill
        cell.font = white_bold_font
        cell.alignment = Alignment(horizontal="left", vertical="center")
        cell.border = thin_border

    key_points = data.get("key_points", [])
    row_num = 2
    for idx, kp in enumerate(key_points, 1):
        ws3.row_dimensions[row_num].height = 22
        point_text = kp.get("point") if isinstance(kp, dict) else str(kp)
        source_pg = kp.get("source_page") if isinstance(kp, dict) else None

        c1 = ws3.cell(row=row_num, column=1, value=idx)
        c2 = ws3.cell(row=row_num, column=2, value=point_text)
        c3 = ws3.cell(row=row_num, column=3, value=f"Page {source_pg}" if source_pg else "—")

        c1.alignment = Alignment(horizontal="center", vertical="center")
        c1.font = field_font
        c2.font = val_font
        c3.font = Font(name="Calibri", size=10, color="2563EB", bold=True)

        for c in (c1, c2, c3):
            c.border = thin_border
            if row_num % 2 == 0:
                c.fill = subtle_row_fill
        row_num += 1

    # ── SHEET 4: Metadata ────────────────────────────────────────────────────
    ws4 = wb.create_sheet(title="Metadata & Warnings")
    ws4.views.sheetView[0].showGridLines = True

    meta_items = [
        ("Original Filename", data.get("original_filename", "")),
        ("Document ID", data.get("document_id", "")),
        ("Extraction Run ID", data.get("run_id", "")),
        ("Config ID", data.get("config_id", "cfg-default-fir")),
        ("Config Version", str(data.get("config_version", 1))),
        ("Extraction Timestamp", data.get("created_at", datetime.now(timezone.utc).isoformat())),
        ("Review Status", data.get("status", "COMPLETED")),
        ("Extraction Warnings", "; ".join(data.get("warnings", [])) or "None"),
    ]
    headers4 = ["Property", "Value"]
    for col_idx, h in enumerate(headers4, 1):
        cell = ws4.cell(row=1, column=col_idx, value=h)
        cell.fill = navy_header_fill
        cell.font = white_bold_font
        cell.border = thin_border

    for row_idx, (k, v) in enumerate(meta_items, 2):
        c1 = ws4.cell(row=row_idx, column=1, value=k)
        c2 = ws4.cell(row=row_idx, column=2, value=v)
        c1.font = field_font
        c2.font = val_font
        c1.border = thin_border
        c2.border = thin_border

    # Auto-fit columns for all sheets
    for ws in (ws1, ws2, ws3, ws4):
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val_s = str(cell.value or "")
                if len(val_s) > max_len:
                    max_len = len(val_s)
            ws.column_dimensions[col_letter].width = min(max(max_len + 4, 14), 70)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output.getvalue()


# ─────────────────────────────────────────────────────────────────────────────
# 2. WORD (.DOCX) EXPORT GENERATOR
# ─────────────────────────────────────────────────────────────────────────────

def _set_cell_background(cell, fill_hex: str):
    """Set background color of a Word table cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)


def _set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Set inner margins (padding) of a Word table cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)


def generate_docx_bytes(data: Dict[str, Any]) -> bytes:
    """
    Generate professional official Word document (.docx):
    - AP Government and Police / Excise header
    - Document metadata summary box
    - Structured FIR Details Table
    - Extracted Entities Table
    - Key Points & Summary
    - Official certification signature block
    """
    doc = docx.Document()

    # Page Margins: 0.75 inch
    for section in doc.sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)

    # ── HEADER ───────────────────────────────────────────────────────────────
    p_gov = doc.add_paragraph()
    p_gov.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_gov = p_gov.add_run("GOVERNMENT OF ANDHRA PRADESH\nPROHIBITION & EXCISE / POLICE DEPARTMENT")
    r_gov.bold = True
    r_gov.font.size = Pt(13)
    r_gov.font.color.rgb = RGBColor(17, 45, 78)  # Navy #112D4E

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_sub = p_sub.add_run("AI-POWERED FIR DOCUMENT EXTRACTION & VERIFICATION RECORD")
    r_sub.bold = True
    r_sub.font.size = Pt(11)
    r_sub.font.color.rgb = RGBColor(15, 81, 50)  # Emerald #0F5132

    p_rule = doc.add_paragraph()
    p_rule.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_rule = p_rule.add_run("—" * 50)
    r_rule.font.size = Pt(9)
    r_rule.font.color.rgb = RGBColor(203, 213, 225)

    # ── METADATA BOX ─────────────────────────────────────────────────────────
    meta_table = doc.add_table(rows=2, cols=4)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_table.autofit = False

    col_widths = [Inches(1.8), Inches(1.8), Inches(1.6), Inches(1.8)]
    for row in meta_table.rows:
        for idx, width in enumerate(col_widths):
            row.cells[idx].width = width

    fir_num = "—"
    station_str = "—"
    for f in data.get("structured_fields", []):
        if f.get("key") == "fir_number" and f.get("value"):
            fir_num = str(f["value"])
        if f.get("key") == "station" and f.get("value"):
            station_str = str(f["value"])

    meta_cells = [
        (0, 0, "FIR / Crime Number:", fir_num),
        (0, 2, "Station Jurisdiction:", station_str),
        (1, 0, "Original File:", data.get("original_filename", "FIR Document")),
        (1, 2, "Extraction Date:", datetime.now().strftime("%d-%b-%Y %H:%M")),
    ]
    for r_idx, c_idx, label, val in meta_cells:
        cell_lbl = meta_table.cell(r_idx, c_idx)
        cell_lbl.text = label
        p_l = cell_lbl.paragraphs[0]
        p_l.runs[0].bold = True
        p_l.runs[0].font.size = Pt(9.5)
        p_l.runs[0].font.color.rgb = RGBColor(71, 85, 105)
        _set_cell_background(cell_lbl, "F1F5F9")
        _set_cell_margins(cell_lbl, top=80, bottom=80, left=100, right=100)

        cell_val = meta_table.cell(r_idx, c_idx + 1)
        cell_val.text = val
        p_v = cell_val.paragraphs[0]
        p_v.runs[0].font.size = Pt(9.5)
        p_v.runs[0].font.color.rgb = RGBColor(15, 23, 42)
        _set_cell_background(cell_val, "FFFFFF")
        _set_cell_margins(cell_val, top=80, bottom=80, left=100, right=100)

    doc.add_paragraph()  # Spacing

    # ── SECTION 1: STRUCTURED FIELDS TABLE ───────────────────────────────────
    h1 = doc.add_paragraph()
    r_h1 = h1.add_run("1. STRUCTURED FIR PARTICULARS")
    r_h1.bold = True
    r_h1.font.size = Pt(11)
    r_h1.font.color.rgb = RGBColor(17, 45, 78)

    fields = data.get("structured_fields", [])
    t_fields = doc.add_table(rows=1, cols=3)
    t_fields.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_fields.autofit = False

    t_widths = [Inches(2.2), Inches(3.8), Inches(1.0)]
    hdr_row = t_fields.rows[0]
    headers = ["Field Name", "Extracted / Verified Value", "Citation"]
    for idx, (h, w) in enumerate(zip(headers, t_widths)):
        c = hdr_row.cells[idx]
        c.width = w
        c.text = h
        p = c.paragraphs[0]
        p.runs[0].bold = True
        p.runs[0].font.size = Pt(9.5)
        p.runs[0].font.color.rgb = RGBColor(255, 255, 255)
        _set_cell_background(c, "112D4E")
        _set_cell_margins(c, top=80, bottom=80, left=100, right=100)

    for r_i, field in enumerate(fields):
        row = t_fields.add_row()
        val = field.get("value", "")
        if isinstance(val, list):
            val = ", ".join(str(v) for v in val)
        elif val is None:
            val = "—"
        else:
            val = str(val)

        vals = [
            field.get("label", field.get("key", "")),
            val,
            f"Page {field.get('source_page')}" if field.get("source_page") else "—"
        ]
        bg = "F8FAFC" if r_i % 2 == 1 else "FFFFFF"

        for idx, (v_text, w) in enumerate(zip(vals, t_widths)):
            c = row.cells[idx]
            c.width = w
            c.text = v_text
            p = c.paragraphs[0]
            if len(p.runs) > 0:
                p.runs[0].font.size = Pt(9)
                if idx == 0:
                    p.runs[0].bold = True
                    p.runs[0].font.color.rgb = RGBColor(30, 41, 59)
                elif idx == 2:
                    p.runs[0].font.color.rgb = RGBColor(37, 99, 235)
                    p.runs[0].bold = True
                else:
                    p.runs[0].font.color.rgb = RGBColor(15, 23, 42)
            _set_cell_background(c, bg)
            _set_cell_margins(c, top=60, bottom=60, left=80, right=80)

    doc.add_paragraph()  # Spacing

    # ── SECTION 2: ENTITIES SUMMARY ──────────────────────────────────────────
    entities = data.get("entities", [])
    if entities:
        h2 = doc.add_paragraph()
        r_h2 = h2.add_run("2. IDENTIFIED NAMED ENTITIES & RECORD LINKAGES")
        r_h2.bold = True
        r_h2.font.size = Pt(11)
        r_h2.font.color.rgb = RGBColor(15, 81, 50)

        t_ent = doc.add_table(rows=1, cols=4)
        t_ent.alignment = WD_TABLE_ALIGNMENT.CENTER
        t_ent.autofit = False

        ent_widths = [Inches(2.5), Inches(1.8), Inches(1.5), Inches(1.2)]
        hdr_ent = t_ent.rows[0]
        ent_hdrs = ["Entity Name", "Classification Type", "Confidence", "Citation"]
        for idx, (h, w) in enumerate(zip(ent_hdrs, ent_widths)):
            c = hdr_ent.cells[idx]
            c.width = w
            c.text = h
            p = c.paragraphs[0]
            p.runs[0].bold = True
            p.runs[0].font.size = Pt(9.5)
            p.runs[0].font.color.rgb = RGBColor(255, 255, 255)
            _set_cell_background(c, "0F5132")
            _set_cell_margins(c, top=80, bottom=80, left=100, right=100)

        for r_i, ent in enumerate(entities):
            row = t_ent.add_row()
            conf = ent.get("confidence")
            conf_str = f"{int(conf * 100)}%" if isinstance(conf, (int, float)) else "95%"
            vals = [
                ent.get("name", ""),
                ent.get("type", "ENTITY"),
                conf_str,
                f"Page {ent.get('source_page', 1)}"
            ]
            bg = "F8FAFC" if r_i % 2 == 1 else "FFFFFF"
            for idx, (v_text, w) in enumerate(zip(vals, ent_widths)):
                c = row.cells[idx]
                c.width = w
                c.text = v_text
                p = c.paragraphs[0]
                if len(p.runs) > 0:
                    p.runs[0].font.size = Pt(9)
                    if idx == 0:
                        p.runs[0].bold = True
                _set_cell_background(c, bg)
                _set_cell_margins(c, top=60, bottom=60, left=80, right=80)

        doc.add_paragraph()  # Spacing

    # ── SECTION 3: KEY FACTS & SUMMARY ───────────────────────────────────────
    key_points = data.get("key_points", [])
    if key_points:
        h3 = doc.add_paragraph()
        r_h3 = h3.add_run("3. FACTUAL KEY POINTS & HIGHLIGHTS")
        r_h3.bold = True
        r_h3.font.size = Pt(11)
        r_h3.font.color.rgb = RGBColor(180, 83, 9)

        for kp in key_points:
            pt_text = kp.get("point") if isinstance(kp, dict) else str(kp)
            src_pg = kp.get("source_page") if isinstance(kp, dict) else None
            p_kp = doc.add_paragraph(style="List Bullet")
            r_pt = p_kp.add_run(pt_text)
            r_pt.font.size = Pt(9.5)
            if src_pg:
                r_cit = p_kp.add_run(f" [Page {src_pg}]")
                r_cit.font.size = Pt(8.5)
                r_cit.font.color.rgb = RGBColor(37, 99, 235)
                r_cit.bold = True

    # ── FOOTER & OFFICER CERTIFICATION ───────────────────────────────────────
    doc.add_paragraph()
    p_cert = doc.add_paragraph()
    r_cert = p_cert.add_run(
        "OFFICER CERTIFICATION: I have personally examined the extracted data against the original uploaded FIR "
        "document. All particulars, accused names, seized property details, and invoked sections are verified as accurate."
    )
    r_cert.italic = True
    r_cert.font.size = Pt(8.5)
    r_cert.font.color.rgb = RGBColor(100, 116, 139)

    p_sig = doc.add_paragraph()
    p_sig.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r_sig = p_sig.add_run("\n\n___________________________________\nSignature of Verifying Officer\nName / Rank / Station Seal")
    r_sig.font.size = Pt(9)
    r_sig.bold = True
    r_sig.font.color.rgb = RGBColor(30, 41, 59)

    output = io.BytesIO()
    doc.save(output)
    output.seek(0)
    return output.getvalue()


# ─────────────────────────────────────────────────────────────────────────────
# 3. CSV (.CSV) EXPORT GENERATOR
# ─────────────────────────────────────────────────────────────────────────────

def generate_csv_bytes(data: Dict[str, Any]) -> bytes:
    """Generate RFC 4180 compliant CSV bytes of structured fields."""
    output = io.StringIO()
    writer = csv.writer(output, quoting=csv.QUOTE_MINIMAL)

    # Header
    writer.writerow([
        "Field Key",
        "Field Label",
        "Data Type",
        "Extracted / Reviewed Value",
        "Source Page Citation",
        "Required Field",
        "Document Name",
        "FIR Number"
    ])

    doc_name = data.get("original_filename", "")
    fir_number = ""
    for f in data.get("structured_fields", []):
        if f.get("key") == "fir_number":
            fir_number = str(f.get("value") or "")
            break

    for field in data.get("structured_fields", []):
        val = field.get("value", "")
        if isinstance(val, list):
            val = ", ".join(str(v) for v in val)
        elif val is None:
            val = ""
        else:
            val = str(val)

        writer.writerow([
            field.get("key", ""),
            field.get("label", ""),
            field.get("data_type", "string"),
            val,
            field.get("source_page") or "",
            "YES" if field.get("required") else "NO",
            doc_name,
            fir_number
        ])

    return output.getvalue().encode("utf-8-sig")  # BOM for Excel CSV compatibility


# ─────────────────────────────────────────────────────────────────────────────
# 4. JSON (.JSON) EXPORT GENERATOR
# ─────────────────────────────────────────────────────────────────────────────

def generate_json_bytes(data: Dict[str, Any]) -> bytes:
    """Generate indented JSON bytes of complete extraction data."""
    export_payload = {
        "$schema": "https://fir-visionai.ap.gov.in/schemas/fir-extraction-v1.json",
        "metadata": {
            "application": "FIR VisionAI — Unified FIR & Case Intelligence Platform",
            "department": "Andhra Pradesh Prohibition & Excise / Police Department",
            "document_id": data.get("document_id"),
            "original_filename": data.get("original_filename"),
            "run_id": data.get("run_id"),
            "config_id": data.get("config_id", "cfg-default-fir"),
            "config_version": data.get("config_version", 1),
            "status": data.get("status", "COMPLETED"),
            "exported_at": datetime.now(timezone.utc).isoformat(),
        },
        "structured_fields": data.get("structured_fields", []),
        "entities": data.get("entities", []),
        "key_points": data.get("key_points", []),
        "warnings": data.get("warnings", []),
        "raw_document_content": data.get("document_content", "")
    }
    return json.dumps(export_payload, indent=2, ensure_ascii=False).encode("utf-8")
