"""
Reusable document ingestion: validate -> extract text -> clean -> chunk -> embed -> index -> verify.

Status rules (Document.ingestion_status / Document.index_status / IngestionJob.status):
- No extractable text at all          -> NEEDS_REVIEW, nothing searchable, no placeholder text stored.
- Text indexed and verified           -> COMPLETED / INDEXED.
- Some pages unreadable, or embeddings
  expected but not produced            -> COMPLETED / PARTIALLY_INDEXED.
- Verification query cannot find it   -> COMPLETED / FAILED (not searchable).
- Invalid/corrupt file or crash       -> FAILED / FAILED with the error recorded on the job.
"""
import asyncio
import json
import logging
import os
import tempfile
from datetime import datetime, timedelta
from typing import List, Optional, Tuple

from sqlalchemy import inspect as sa_inspect
from sqlalchemy import select, text
from sqlalchemy.ext.asyncio import AsyncSession

from ..config import settings
from ..models.base import generate_uuid
from ..models.document import Document, DocumentChunk, DocumentPage, IngestionJob
from .embeddings import embedding_service, to_json, to_pgvector
from .text_processing import (
    PageText,
    build_chunks,
    clean_extracted_text,
    section_label_for,
    segment_pages,
)

logger = logging.getLogger(__name__)

MIN_PAGE_TEXT = 20  # characters; below this a page is treated as unreadable
STALE_PROCESSING_AFTER = timedelta(minutes=15)
SEARCHABLE_INDEX_STATES = ("INDEXED", "PARTIALLY_INDEXED")


class IngestionError(Exception):
    """A file that cannot be ingested (corrupt, empty, unsupported)."""


# ── Text extraction ──────────────────────────────────────────────────────────

PDF_RENDER_DPI = 300        # scanned PDF pages are rasterised at this resolution for OCR
PDF_RETRY_DPI = 400         # and once more at this one when the first read is too short
MIN_OCR_EDGE = 1800         # small phone photos are upscaled to this long edge before OCR


def enhance_for_ocr(path: str) -> Optional[str]:
    """
    A cleaned-up copy of a photo or scan for a second OCR attempt: EXIF rotation applied, greyscale,
    straightened (deskew by the angle with the sharpest row profile), contrast stretched, upscaled
    when small and sharpened. Returns a temporary PNG path, or None if the image cannot be processed.
    """
    try:
        from PIL import Image, ImageFilter, ImageOps
        with Image.open(path) as src:
            img = ImageOps.exif_transpose(src).convert("L")
        long_edge = max(img.size)
        if long_edge < MIN_OCR_EDGE:
            scale = MIN_OCR_EDGE / long_edge
            img = img.resize((int(img.width * scale), int(img.height * scale)), Image.LANCZOS)
        img = ImageOps.autocontrast(img, cutoff=1)
        img = img.rotate(_skew_angle(img), resample=Image.BICUBIC, expand=True, fillcolor=255)
        img = img.filter(ImageFilter.SHARPEN)
        with tempfile.NamedTemporaryFile(suffix=".png", delete=False) as tmp:
            img.save(tmp.name, format="PNG")
            return tmp.name
    except Exception as e:
        logger.debug(f"Image enhancement for OCR failed: {e}")
        return None


def _skew_angle(img, max_angle: float = 6.0, step: float = 0.5) -> float:
    """Rotation (degrees) that best aligns text lines horizontally; 0 when it cannot be measured."""
    try:
        import numpy as np
        small = img.copy()
        small.thumbnail((800, 800))
        best, best_score = 0.0, -1.0
        for i in range(int(-max_angle / step), int(max_angle / step) + 1):
            angle = i * step
            arr = np.asarray(small.rotate(angle, expand=False, fillcolor=255), dtype=np.float32)
            score = float(np.var((arr < 128).sum(axis=1)))  # text rows vs gaps are sharpest when level
            if score > best_score:
                best, best_score = angle, score
        return best
    except Exception:
        return 0.0


def _ocr_image_file(path: str, mime: str) -> str:
    """
    OCR one image with the primary engine (which retries transient failures with backoff); when the
    read is empty or too short, OCR an enhanced copy (deskewed, contrast-stretched, upscaled).
    Placeholder notices ("no visible text") count as no text.
    """
    from .gemini import gemini_service
    from .text_processing import is_placeholder_text
    text_ = (gemini_service.extract_text_from_image(path, mime) or "").strip()
    if len(text_) >= MIN_PAGE_TEXT * 5 and not is_placeholder_text(text_):
        return text_
    enhanced = enhance_for_ocr(path)
    if enhanced:
        try:
            second = (gemini_service.extract_text_from_image(enhanced, "image/png") or "").strip()
        finally:
            os.remove(enhanced)
        if len(second) > len(text_) and not is_placeholder_text(second):
            text_ = second
    return "" if is_placeholder_text(text_) else text_


def _render_pdf_page(path: str, idx: int, dpi: int) -> str:
    import pypdfium2
    pdf = pypdfium2.PdfDocument(path)
    try:
        img = pdf[idx].render(scale=dpi / 72).to_pil()
    finally:
        pdf.close()
    with tempfile.NamedTemporaryFile(suffix=".png", delete=False) as tmp:
        img.save(tmp.name, format="PNG")
        return tmp.name


def _extract_pdf(path: str, warnings: List[str]) -> List[PageText]:
    from pypdf import PdfReader
    try:
        reader = PdfReader(path)
        if reader.is_encrypted and not reader.decrypt(""):
            raise IngestionError("The PDF is password-protected and cannot be read.")
        page_count = len(reader.pages)
    except IngestionError:
        raise
    except Exception as e:
        raise IngestionError(f"Invalid or corrupted PDF: {str(e)[:200]}")
    if page_count == 0:
        raise IngestionError("The PDF contains no pages.")

    plumber = None
    pages: List[PageText] = []
    for idx, page in enumerate(reader.pages):
        try:
            raw = page.extract_text() or ""
        except Exception as e:
            warnings.append(f"Page {idx + 1}: text layer unreadable ({str(e)[:80]})")
            raw = ""
        ocr_applied = False
        if len(raw.strip()) < MIN_PAGE_TEXT:
            try:
                import pdfplumber
                plumber = plumber or pdfplumber.open(path)
                raw = plumber.pages[idx].extract_text() or raw
            except Exception:
                pass
        if len(raw.strip()) < MIN_PAGE_TEXT:
            # Scanned page: rasterise at 300 DPI and OCR, then once more at 400 DPI if the read is short
            for dpi in (PDF_RENDER_DPI, PDF_RETRY_DPI):
                try:
                    tmp_name = _render_pdf_page(path, idx, dpi)
                    try:
                        ocr = _ocr_image_file(tmp_name, "image/png")
                    finally:
                        os.remove(tmp_name)
                    if len(ocr) > len(raw.strip()) and len(ocr) >= MIN_PAGE_TEXT:
                        raw, ocr_applied = ocr, True
                    if len(ocr) >= MIN_PAGE_TEXT * 5:
                        break
                except Exception as e:
                    logger.debug(f"OCR for page {idx + 1} at {dpi} DPI unavailable: {e}")
        cleaned = clean_extracted_text(raw)
        unreadable = len(cleaned) < MIN_PAGE_TEXT
        if unreadable:
            warnings.append(f"Page {idx + 1}: no machine-readable text (OCR unavailable or failed)")
        pages.append(PageText(idx + 1, "" if unreadable else cleaned, ocr_applied, unreadable))
    if plumber:
        plumber.close()
    return pages


def _extract_docx(path: str) -> List[PageText]:
    try:
        import docx
        d = docx.Document(path)
    except Exception as e:
        raise IngestionError(f"Invalid or corrupted DOCX: {str(e)[:200]}")
    parts = [p.text for p in d.paragraphs if p.text.strip()]
    for table in d.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    # DOCX has no fixed pagination; the whole text is page 1 and citations rely on sections
    text_ = clean_extracted_text("\n".join(parts))
    return [PageText(1, text_, False, len(text_) < MIN_PAGE_TEXT)]


async def _extract_image(path: str, mime: str, warnings: List[str]) -> List[PageText]:
    try:
        from PIL import Image, ImageOps
        with Image.open(path) as img:
            ImageOps.exif_transpose(img).load()
    except Exception as e:
        raise IngestionError(f"Corrupted or invalid image file: {str(e)[:200]}")
    text_ = await asyncio.to_thread(_ocr_image_file, path, mime)
    if len(text_) < MIN_PAGE_TEXT * 5:
        # Second OCR engine
        second = (await _second_engine_pages(path, mime, [1], warnings)).get(1, "")
        if len(second) > len(text_):
            text_ = second
    cleaned = clean_extracted_text(text_)
    unreadable = len(cleaned) < MIN_PAGE_TEXT
    if unreadable:
        warnings.append("Image: no text extracted (OCR provider unavailable or image unreadable)")
    return [PageText(1, "" if unreadable else cleaned, True, unreadable)]


def unreadable_reason(warnings: List[str]) -> str:
    """A specific, actionable reason for a document from which no text could be read."""
    from .gemini import gemini_service
    joined = " ".join(warnings).lower()
    provider = getattr(gemini_service, "_last_failure_reason", None)
    if not gemini_service.client:
        return "OCR_UNAVAILABLE: no OCR provider is configured (GEMINI_API_KEY / SARVAM_API_KEY). Configure one and press Re-process."
    if provider == "QUOTA_EXHAUSTED":
        return "OCR_QUOTA: the OCR provider's quota is used up. Press Re-process after the quota resets."
    if provider in ("PROVIDER_TIMEOUT", "PROVIDER_OVERLOADED", "PROVIDER_UNREACHABLE"):
        return f"OCR_PROVIDER_FAILED ({provider}): OCR was retried with backoff and still failed. Press Re-process later."
    if "sarvam" in joined and "failed" in joined:
        return "OCR_NO_TEXT: neither OCR engine found text on the page (the second engine failed). Check the file and re-upload a clearer scan."
    return ("OCR_NO_TEXT: both OCR attempts (original and enhanced/deskewed image) found no readable text. "
            "The file may be blank, a photograph without text, or too blurred: re-upload a clearer scan or archive it.")


async def _second_engine_pages(path: str, mime: str, pages: List[int], warnings: List[str]) -> dict:
    """Page text from the second OCR engine (Sarvam Digitise), or {} when it is not configured or fails."""
    from .text_processing import is_placeholder_text
    try:
        from .sarvam_doc_ai import sarvam_doc_ai
        if not sarvam_doc_ai.is_available():
            return {}
        res = await sarvam_doc_ai.digitise_pages(path, mime, only_pages=pages)
        return {p: t.strip() for p, t in res.items() if t and not is_placeholder_text(t)}
    except Exception as e:
        warnings.append(f"Second OCR engine (Sarvam Digitise) failed: {str(e)[:120]}")
        return {}


async def extract_pages(path: str, mime: str, filename: str = "") -> Tuple[List[PageText], List[str]]:
    """Extract per-page text. Raises IngestionError for files that cannot be ingested."""
    warnings: List[str] = []
    if not os.path.exists(path) or os.path.getsize(path) == 0:
        raise IngestionError("The file is empty or missing.")
    name = filename.lower()
    if mime == "application/pdf" or name.endswith(".pdf"):
        pages = await asyncio.to_thread(_extract_pdf, path, warnings)
        missing = [p.page_number for p in pages if p.unreadable]
        if missing:
            # Pages the primary OCR could not read go to the second engine (in 10-page batches)
            recovered = await _second_engine_pages(path, "application/pdf", missing, warnings)
            for p in pages:
                text_ = clean_extracted_text(recovered.get(p.page_number, ""))
                if p.unreadable and len(text_) >= MIN_PAGE_TEXT:
                    p.text, p.ocr_applied, p.unreadable = text_, True, False
                    warnings.append(f"Page {p.page_number}: read by the second OCR engine")
    elif mime.startswith("image/"):
        pages = await _extract_image(path, mime, warnings)
    elif name.endswith(".docx") or "wordprocessingml" in mime:
        pages = await asyncio.to_thread(_extract_docx, path)
    elif name.endswith((".txt", ".md")) or mime.startswith("text/"):
        with open(path, "rb") as f:
            body = clean_extracted_text(f.read().decode("utf-8", errors="replace"))
        pages = [PageText(1, body, False, len(body) < MIN_PAGE_TEXT)]
    else:
        raise IngestionError(f"Unsupported document type '{mime or name}'.")
    return pages, warnings


# ── Index storage helpers ────────────────────────────────────────────────────

def _is_sqlite(db: AsyncSession) -> bool:
    try:
        return db.get_bind().dialect.name == "sqlite"
    except Exception:
        return False


_pg_vector_column: Optional[bool] = None


async def has_pg_vector_column(db: AsyncSession) -> bool:
    """True when the PostgreSQL chunk table has the pgvector 'embedding' column (added by migration)."""
    global _pg_vector_column
    if _is_sqlite(db):
        return False
    if _pg_vector_column is None:
        def _check(sync_session):
            cols = sa_inspect(sync_session.connection()).get_columns("document_chunks")
            return any(c["name"] == "embedding" for c in cols)
        try:
            _pg_vector_column = await db.run_sync(_check)
        except Exception as e:
            # Not cached: a transient connection error must not switch vectors to the JSON column for good
            logger.warning(f"Could not inspect document_chunks for the pgvector column: {type(e).__name__}: {e}")
            return False
    return _pg_vector_column


async def insert_fts_row(db: AsyncSession, chunk_id: str, document_id: str, page: int, chunk_text: str) -> None:
    """SQLite FTS5 row for a chunk; PostgreSQL uses the generated tsvector column instead."""
    if not _is_sqlite(db):
        return
    try:
        await db.execute(
            text("INSERT INTO document_chunks_fts (chunk_id, document_id, physical_page_number, chunk_text) "
                 "VALUES (:c, :d, :p, :t)"),
            {"c": chunk_id, "d": document_id, "p": page, "t": chunk_text},
        )
    except Exception as e:
        logger.warning(f"FTS5 insert failed: {e}")


async def delete_index_rows(db: AsyncSession, document_id: str, keep_pages: bool = False) -> None:
    await db.execute(text("DELETE FROM document_chunks WHERE document_id = :d"), {"d": document_id})
    if not keep_pages:
        await db.execute(text("DELETE FROM document_pages WHERE document_id = :d"), {"d": document_id})
    if _is_sqlite(db):
        try:
            await db.execute(text("DELETE FROM document_chunks_fts WHERE document_id = :d"), {"d": document_id})
        except Exception:
            pass


# ── Indexing ─────────────────────────────────────────────────────────────────

async def index_document(
    db: AsyncSession,
    doc: Document,
    pages: List[PageText],
    *,
    trigger: str,
    user_id: Optional[str],
    warnings: Optional[List[str]] = None,
    job: Optional[IngestionJob] = None,
    rewrite_pages: bool = True,
) -> IngestionJob:
    """Persist pages and chunks, embed, verify, and set statuses. Commits; never raises for data issues."""
    warnings = list(warnings or [])
    if job is None:
        job = IngestionJob(id=generate_uuid(), document_id=doc.id, trigger=trigger, created_by=user_id,
                           started_at=datetime.utcnow())
        db.add(job)
    job.status, job.stage = "PROCESSING", "CHUNK"
    doc.ingestion_status = "PROCESSING"
    await db.flush()

    for p in pages:
        # Placeholder or "no visible text" notices are never stored or indexed as page text
        if p.text and _is_legacy_placeholder(p.text):
            p.text, p.unreadable = "", True
    await delete_index_rows(db, doc.id, keep_pages=not rewrite_pages)
    if rewrite_pages:
        for p in pages:
            db.add(DocumentPage(
                id=generate_uuid(), document_id=doc.id, physical_page_number=p.page_number,
                printed_page_number=p.page_number, raw_text=p.text or None, ocr_applied=p.ocr_applied,
                ocr_confidence=None if p.unreadable else (0.9 if p.ocr_applied else 1.0),
                is_unreadable=p.unreadable,
            ))
    doc.physical_page_count = len(pages)
    doc.is_scanned = any(p.ocr_applied for p in pages)

    readable = [p for p in pages if not p.unreadable and p.text.strip()]
    job.pages_total, job.pages_with_text = len(pages), len(readable)

    if not readable:
        doc.ingestion_status = "NEEDS_REVIEW"
        doc.index_status = "NEEDS_REVIEW"
        doc.indexed_chunk_count = doc.embedded_chunk_count = 0
        doc.error_message = unreadable_reason(warnings)
        job.status, job.stage, job.error_message = "NEEDS_REVIEW", "EXTRACT", doc.error_message
        return await _finish(db, doc, job, warnings)

    segments, _ = segment_pages(readable, section_label_for(doc.document_title or doc.original_filename, doc.doc_category))
    chunks = build_chunks(segments)
    chunk_ids: List[str] = []
    for c in chunks:
        cid = generate_uuid()
        chunk_ids.append(cid)
        db.add(DocumentChunk(
            id=cid, document_id=doc.id, physical_page_number=c["physical_page_number"],
            chunk_index=c["chunk_index"], chunk_text=c["chunk_text"], char_start=0,
            char_end=len(c["chunk_text"]),
            # Column limits (PostgreSQL enforces them; SQLite does not)
            section_ref=(c["section_ref"] or "")[:100] or None,
            section_title=(c["section_title"] or "")[:300] or None,
        ))
    await db.flush()
    for cid, c in zip(chunk_ids, chunks):
        # Index the section reference/title with the text so "theft" finds "Section 303 (Of theft)"
        header = " ".join(x for x in (c["section_ref"], c["section_title"]) if x)
        await insert_fts_row(db, cid, doc.id, c["physical_page_number"],
                             f"{header}\n{c['chunk_text']}" if header else c["chunk_text"])
    job.chunks_written = len(chunks)

    # Embeddings (semantic search). Absent provider => keyword-only, reported as such.
    job.stage = "EMBED"
    embedded = 0
    if embedding_service.is_available():
        vectors = await asyncio.to_thread(embedding_service.embed_documents, [c["chunk_text"] for c in chunks])
        if vectors:
            use_pg = await has_pg_vector_column(db)
            for cid, vec in zip(chunk_ids, vectors):
                if use_pg:
                    await db.execute(text("UPDATE document_chunks SET embedding = CAST(:v AS extensions.vector) WHERE id = :id"),
                                     {"v": to_pgvector(vec), "id": cid})
                else:
                    await db.execute(text("UPDATE document_chunks SET embedding_json = :v WHERE id = :id"),
                                     {"v": to_json(vec), "id": cid})
            embedded = len(vectors)
            doc.embedding_model = embedding_service.model
        else:
            reason = getattr(embedding_service, "last_error", None)
            warnings.append("Embedding generation failed" + (f" ({reason})" if reason else "")
                            + "; document is searchable by keywords only.")
    else:
        warnings.append("No embedding provider configured; document is searchable by keywords only.")
    job.chunks_embedded = embedded
    await db.flush()

    # Verification: the index must actually return this document for a query built from its text
    job.stage = "VERIFY"
    verified = await verify_document_searchable(db, doc.id, chunks)

    doc.ingestion_status = "COMPLETED"
    doc.indexed_chunk_count = len(chunks)
    doc.embedded_chunk_count = embedded
    doc.error_message = None
    if not verified:
        doc.index_status = "FAILED"
        doc.error_message = "Index verification failed: a test search did not return this document."
        job.status, job.error_message = "FAILED", doc.error_message
    elif len(readable) < len(pages) or (embedding_service.is_available() and embedded < len(chunks)):
        doc.index_status = "PARTIALLY_INDEXED"
        job.status = "PARTIALLY_INDEXED"
    else:
        doc.index_status = "INDEXED"
        job.status = "INDEXED"
    doc.indexed_at = datetime.utcnow()
    return await _finish(db, doc, job, warnings)


async def _finish(db: AsyncSession, doc: Document, job: IngestionJob, warnings: List[str]) -> IngestionJob:
    job.warnings_json = json.dumps(warnings[:50]) if warnings else None
    job.finished_at = datetime.utcnow()
    await db.commit()
    await db.refresh(doc)
    await db.refresh(job)
    try:
        from .rag_engine import rag_engine
        rag_engine.invalidate_cache(doc.id)
    except Exception:
        pass
    return job


async def verify_document_searchable(db: AsyncSession, document_id: str, chunks: List[dict]) -> bool:
    """Run a real keyword search for distinctive words from the document, restricted to it."""
    from .retrieval import keyword_search, probe_terms
    sample = max(chunks, key=lambda c: len(c["chunk_text"]))["chunk_text"] if chunks else ""
    terms = probe_terms(sample)
    if not terms:
        return bool(chunks)  # text with no indexable words (e.g. digits only) is stored but unverifiable
    hits = await keyword_search(db, " ".join(terms), document_ids=[document_id], limit=3)
    return any(h["document_id"] == document_id for h in hits)


async def mark_failed(db: AsyncSession, doc_id: str, job_id: Optional[str], error: str) -> IngestionJob:
    """Record a failure in a fresh transaction. Takes ids because rollback expires loaded objects."""
    try:
        await db.rollback()
    except Exception:
        pass
    doc = (await db.execute(select(Document).where(Document.id == doc_id))).scalar_one()
    job = None
    if job_id:
        job = (await db.execute(select(IngestionJob).where(IngestionJob.id == job_id))).scalar_one_or_none()
    if job is None:
        job = IngestionJob(id=generate_uuid(), document_id=doc_id, trigger="UPLOAD", started_at=datetime.utcnow())
        db.add(job)
    doc.ingestion_status = "FAILED"
    doc.index_status = "FAILED"
    doc.error_message = error[:500]
    job.status, job.error_message = "FAILED", error[:2000]
    return await _finish(db, doc, job, [])


async def ingest_file(
    db: AsyncSession,
    doc: Document,
    file_path: str,
    *,
    trigger: str,
    user_id: Optional[str],
) -> IngestionJob:
    """Full pipeline for one stored file. Always returns the job; job.status says what happened."""
    doc_id, job_id = doc.id, generate_uuid()
    job = IngestionJob(id=job_id, document_id=doc_id, trigger=trigger, created_by=user_id,
                       status="PROCESSING", stage="VALIDATE", started_at=datetime.utcnow())
    db.add(job)
    doc.ingestion_status = "PROCESSING"
    await db.commit()
    try:
        job.stage = "EXTRACT"
        pages, warnings = await extract_pages(file_path, doc.mime_type or "", doc.original_filename or "")
        return await index_document(db, doc, pages, trigger=trigger, user_id=user_id, warnings=warnings, job=job)
    except IngestionError as e:
        return await mark_failed(db, doc_id, job_id, str(e))
    except Exception as e:
        logger.error(f"Ingestion crashed for document {doc_id}: {e}", exc_info=True)
        return await mark_failed(db, doc_id, job_id, f"Unexpected ingestion error: {str(e)[:300]}")


async def reindex_from_pages(db: AsyncSession, doc: Document, *, trigger: str = "REINDEX",
                             user_id: Optional[str] = None) -> IngestionJob:
    """Rebuild chunks/embeddings from already-extracted page text (no OCR or file access needed)."""
    rows = (await db.execute(
        select(DocumentPage).where(DocumentPage.document_id == doc.id).order_by(DocumentPage.physical_page_number)
    )).scalars().all()
    pages = [PageText(p.physical_page_number, clean_extracted_text(p.raw_text or ""), bool(p.ocr_applied),
                      bool(p.is_unreadable) or len((p.raw_text or "").strip()) < MIN_PAGE_TEXT) for p in rows]
    for p in pages:
        if _is_legacy_placeholder(p.text):
            p.text, p.unreadable = "", True
    return await index_document(db, doc, pages, trigger=trigger, user_id=user_id, rewrite_pages=False)


async def reindex_after_recovered_text(db: AsyncSession, doc: Document, user_id: Optional[str] = None) -> None:
    """
    Re-index a document whose page text arrived after ingestion (e.g. Sarvam Digitise read a page
    that ingestion's OCR could not). Without this the document stays NEEDS_REVIEW and unsearchable
    although its text exists. Failures are logged; the extraction result is kept either way.
    """
    try:
        await reindex_from_pages(db, doc, trigger="OCR_RECOVERED", user_id=user_id)
    except Exception as e:
        logger.warning(f"Re-indexing {doc.id} after recovered OCR text failed: {type(e).__name__}: {e}")
        try:
            await db.rollback()
        except Exception:
            pass


def _is_legacy_placeholder(t: str) -> bool:
    """Placeholder strings older pipeline versions stored instead of real OCR text."""
    from .text_processing import is_placeholder_text
    return bool(t) and is_placeholder_text(t)


# ── File access for retries ──────────────────────────────────────────────────

async def materialize_document_file(doc: Document) -> Tuple[Optional[str], bool]:
    """Return (local_path, is_temporary) for the document's original file, or (None, False)."""
    candidates = [doc.stored_path or ""]
    fname, fhash = doc.original_filename or "", doc.file_hash or ""
    ev = settings.EVIDENCE_DIR
    candidates += [os.path.join(ev, fname), os.path.join(ev, "cloud_cache", fname)]
    if fhash:
        candidates += [os.path.join(ev, f"{fhash[:12]}_{fname}"), os.path.join(ev, f"{fhash[:8]}_{fname}"),
                       os.path.join(ev, "cloud_cache", f"{fhash[:8]}_{fname}")]
    for c in candidates:
        if c and os.path.isfile(c):
            return c, False
    from .storage import storage_service
    if storage_service.is_cloud_configured() and fhash:
        # A stored local path (older uploads recorded their temporary cache copy) is not a cloud key;
        # cloud uploads are saved as evidence/{hash8}_{filename}
        stored_key = doc.stored_path if doc.stored_path and not os.path.isabs(doc.stored_path) else None
        # Older uploads were also copied to the bucket root under the local file's name ({hash12}_{filename})
        local_name = os.path.basename((doc.stored_path or "").replace("\\", "/")) if doc.stored_path else None
        for cloud_path in dict.fromkeys(filter(None, (stored_key, f"evidence/{fhash[:8]}_{fname}",
                                                      local_name, f"{fhash[:12]}_{fname}"))):
            try:
                data = await storage_service.get_file_bytes(cloud_path, storage_type="SUPABASE")
            except Exception as e:
                logger.info(f"Storage object {cloud_path!r} not readable for document {doc.id}: {e}")
                data = None
            if data:
                suffix = os.path.splitext(fname)[1] or ".bin"
                with tempfile.NamedTemporaryFile(suffix=suffix, delete=False) as tmp:
                    tmp.write(data)
                    return tmp.name, True
    return None, False


def effective_status(doc: Document) -> str:
    """Ingestion status for display: PROCESSING records abandoned by a crash show as STALLED."""
    status_ = doc.ingestion_status or "PENDING"
    if status_ == "PROCESSING" and doc.updated_at and datetime.utcnow() - doc.updated_at > STALE_PROCESSING_AFTER:
        return "STALLED"
    return status_
